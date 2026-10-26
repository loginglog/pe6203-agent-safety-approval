"""Build the English report from its Markdown source using bundled python-docx.

No experiment or model request is launched. Render with the documents skill's
packaged renderer and inspect every page before delivering the DOCX.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "report/report_draft.md"
OUTPUT = ROOT / "report/PE6203_Report_Draft.docx"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def format_table(table, headers):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    columns = len(headers)
    if columns == 2:
        widths = [2.15, 4.75]
    elif columns == 3:
        widths = [1.35, 2.775, 2.775]
    elif columns == 4:
        widths = [2.45, 1.48, 1.48, 1.49] if "Strategy" in headers[0] else [0.48, 2.15, 2.85, 1.42]
    elif columns == 6:
        widths = [1.25, 1.05, 0.85, 1.15, 1.15, 1.45]
    else:
        widths = [6.9 / columns] * columns
    for grid, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid.set(qn("w:w"), str(round(width * 1440)))
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            border.set(qn("w:" + key), value)
        borders.append(border)
    props.append(borders)
    for ri, row in enumerate(table.rows):
        cannot_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cannot_split)
        if ri == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margin = OxmlElement("w:tcMar")
            for edge in ("top", "bottom", "left", "right"):
                element = OxmlElement("w:" + edge)
                element.set(qn("w:w"), "85" if edge in ("top", "bottom") else "100")
                element.set(qn("w:type"), "dxa")
                margin.append(element)
            cell._tc.get_or_add_tcPr().append(margin)
            shade(cell, "263C56" if ri == 0 else "F1F4F7" if ri % 2 == 0 else "FFFFFF")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.keep_with_next = ri < len(table.rows) - 1
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 or columns == 2 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.5 if columns >= 6 else 10)
                    run.font.bold = ri == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if ri == 0 else "000000")


def main():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    section.footer_distance = Inches(0.35)
    styles = document.styles
    for node in styles.element.xpath(".//w:pBdr"):
        node.getparent().remove(node)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        styles[name].font.name = "Arial"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        props = styles[name].element.get_or_add_rPr()
        for tag in ("color",):
            for node in props.findall(qn("w:" + tag)):
                for key in ("themeColor", "themeShade", "themeTint"):
                    node.attrib.pop(qn("w:" + key), None)
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.line_spacing = 1.1
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Title"].font.size = Pt(22)
    styles["Title"].paragraph_format.space_after = Pt(6)
    styles["Subtitle"].font.size = Pt(12)
    styles["Subtitle"].paragraph_format.space_after = Pt(8)
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.size = Pt(11.5)
    for name in ("Heading 1", "Heading 2"):
        styles[name].paragraph_format.space_before = Pt(11)
        styles[name].paragraph_format.space_after = Pt(5)
        styles[name].paragraph_format.keep_with_next = True
    styles["Caption"].font.size = Pt(10)
    styles["Caption"].font.italic = False
    styles["Caption"].font.bold = False
    styles["Caption"].paragraph_format.space_after = Pt(7)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run()
    run.font.name, run.font.size = "Arial", Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = "Safety Approval Policies for Agent Harnesses"
    document.core_properties.subject = "Exploratory AgentDojo Workspace evaluation"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    lines = SOURCE.read_text().splitlines()
    i, in_code = 0, False
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.keep_together = True
            p.paragraph_format.keep_with_next = i + 1 < len(lines) and not lines[i + 1].startswith("```")
            run = p.add_run(line)
            run.font.name, run.font.size = "Courier New", Pt(8.5)
            i += 1
            continue
        if line.startswith("# "):
            document.add_paragraph(line[2:], style="Title")
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [v.strip() for v in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", v) for v in cells):
                    rows.append(cells)
                i += 1
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for row, values in zip(table.rows, rows):
                for cell, value in zip(row.cells, values):
                    cell.text = value
            format_table(table, rows[0])
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = Pt(4)
            continue
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            run = p.add_run()
            # Keep the review-load figure near its analysis without a large gap.
            figure_width = 6.2 if match[2].endswith("hitl_all_tasks.png") else 6.9
            shape = run.add_picture(str(SOURCE.parent / match[2]), width=Inches(figure_width))
            shape._inline.docPr.set("descr", match[1])
        elif line.startswith("Figure ") or line.startswith("Table "):
            p = document.add_paragraph(line, style="Caption")
            p.paragraph_format.keep_with_next = line.startswith("Table ")
        else:
            p = document.add_paragraph(line)
            if line.startswith("The following text is reviewer-v1"):
                p.paragraph_format.keep_with_next = True
            if line == "An Exploratory Evaluation in AgentDojo Workspace":
                p.style = styles["Subtitle"]
            elif line.startswith("PE6203 Group Project"):
                for run in p.runs:
                    run.font.size = Pt(10)
            elif line.startswith("[") and re.match(r"\[\d+\]", line):
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(9.5)
        i += 1
    document.save(OUTPUT)
    print(OUTPUT)
    print("Source words:", len(re.findall(r"\b[\w'-]+\b", SOURCE.read_text())))


if __name__ == "__main__":
    main()
